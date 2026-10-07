"""Builds the AI test report workbook from eval/results/latest.json (written by eval/run.ts).

Run (from covert_call/eval/):  npm run report   (or: python report.py)
Output: eval/results/QuickBite_AI_Test_Report.xlsx

Sheets: Dashboard (KPIs, targets, charts), Test Cases, Results, Consistency, Detection, Checks, Latency,
Performance, Findings, Transcripts.

Two verdicts per check: "strict" (keyword match on Mia's tool calls, the raw signal) and "judged" (an AI judge
reading the dashboard data: does it convey the fact in any wording). Both are shown; neither replaces the other.
"""
import json
import statistics
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
data = json.loads((HERE / "results" / "latest.json").read_text(encoding="utf-8"))
results = data["results"]
OUT = HERE / "results" / "QuickBite_AI_Test_Report.xlsx"

INK, MUTED, BRAND = "1F2937", "6B7280", "E4572E"
PASS_FILL, FAIL_FILL, WARN_FILL = PatternFill("solid", fgColor="DCFCE7"), PatternFill("solid", fgColor="FEE2E2"), PatternFill("solid", fgColor="FEF3C7")
HEAD_FILL, CARD_FILL = PatternFill("solid", fgColor=INK), PatternFill("solid", fgColor="F3F4F6")
THIN = Border(*(Side(style="thin", color="E5E7EB"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")

# Stable test-case numbers per scenario, in run order.
scen_ids = list(dict.fromkeys(r["id"] for r in results))
TC = {sid: f"TC-{i:02d}" for i, sid in enumerate(scen_ids, 1)}
for r in results:
    r.setdefault("run", 1)
    r.setdefault("attempts", 1)
    for c in r["checks"]:
        c.setdefault("semantic", c["pass"])


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def vals(key, src="metrics", rs=None):
    return [r[src][key] for r in (rs or results) if r.get(src) and num(r[src].get(key)) is not None]


def post(key):
    return [r["metrics"]["postCall"][key] for r in results if r["metrics"].get("postCall") and num(r["metrics"]["postCall"].get(key)) is not None]


def stat(xs, fn):
    return round(fn(xs), 1) if xs else None


def pctl(xs, p):
    if not xs:
        return None
    s = sorted(xs)
    return s[min(len(s) - 1, int(p / 100 * len(s)))]


def status(r, key="pass"):
    if r.get("error"):
        return "ERROR"
    ok = sum(bool(c.get(key)) for c in r["checks"])
    return "PASS" if ok == len(r["checks"]) else "PARTIAL" if ok else "FAIL"


def table(ws, headers, rows, widths=None, start_row=1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, Alignment(vertical="center", wrap_text=True), THIN
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment, cell.border = WRAP, THIN
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + max(1, len(rows))}"


def color_words(ws, col, first, last):
    rng = f"{col}{first}:{col}{max(first, last)}"
    for word, fill in (("PASS", PASS_FILL), ("MET", PASS_FILL), ("PARTIAL", WARN_FILL), ("FAIL", FAIL_FILL), ("MISSED", FAIL_FILL), ("ERROR", FAIL_FILL)):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{word}"'], fill=fill))


all_checks = [c for r in results for c in r["checks"]]
strict_ok = sum(c["pass"] for c in all_checks)
judged_ok = sum(bool(c["semantic"]) for c in all_checks)
responses = [t["firstAudioMs"] for r in results for t in r["turnTimings"] if t["turn"] > 0 and num(t["firstAudioMs"]) is not None]
runs = max(r["run"] for r in results)
by_scen = defaultdict(list)
for r in results:
    by_scen[r["id"]].append(r)
reliable = sum(all(status(r, "semantic") == "PASS" for r in rs) for rs in by_scen.values())

# Targets for a live safety call (demo-grade, not production SLOs).
TARGETS = [
    ("Incident on dashboard after call start", "ms", 2000, stat(vals("incidentCreateMs"), statistics.mean), "avg"),
    ("Mia response, p95 (to first audio)", "ms", 2500, pctl(responses, 95), "p95"),
    ("Mia response, median", "ms", 1500, statistics.median(responses) if responses else None, "p50"),
    ("First danger alert after first caller line", "s", 15, stat(vals("firstDangerS"), statistics.mean), "avg"),
    ("Field update reaches Firestore (p95)", "ms", 1500, pctl(vals("writeP95Ms"), 95), "p95"),
    ("Case summary written after hang-up", "ms", 10000, stat(post("consolidateMs"), statistics.mean), "avg"),
    ("Real route guidance answer", "ms", 5000, stat(vals("routeGuidanceMs"), statistics.mean), "avg"),
]

wb = Workbook()

# ---------------------------------------------------------------- Dashboard
ws = wb.active
ws.title = "Dashboard"
ws.sheet_view.showGridLines = False
ws["B2"] = "QuickBite — AI Test Report"
ws["B2"].font = Font(size=20, bold=True, color=INK)
ws["B3"] = (f"Run {data['ranAt'][:19].replace('T', ' ')} UTC · {len(scen_ids)} scenarios × {runs} run(s) = {len(results)} calls · "
            f"Live model {data['liveModel']} · caller {data['callerModel'].split(' → ')[0]} (+fallbacks) · incidents in '{data.get('collection') or 'not recorded'}'")
ws["B3"].font = Font(color=MUTED)
ws["B4"] = "Strict = keyword match on Mia's tool calls. Judged = AI judge reading the dashboard data (same meaning, any wording)."
ws["B4"].font = Font(color=MUTED, italic=True, size=9)

kpis = [
    ("Test calls", len(results), f"{len(scen_ids)} scenarios × {runs}"),
    ("Checks passed (judged)", f"{judged_ok}/{len(all_checks)}", f"{round(100 * judged_ok / max(1, len(all_checks)))}%"),
    ("Checks passed (strict)", f"{strict_ok}/{len(all_checks)}", f"{round(100 * strict_ok / max(1, len(all_checks)))}%"),
    ("Scenarios passing every run", f"{reliable}/{len(scen_ids)}", "judged"),
    ("Incident created (avg)", stat(vals("incidentCreateMs"), statistics.mean), "ms"),
    ("Mia response median", statistics.median(responses) if responses else None, "ms to first audio"),
    ("Mia response p95", pctl(responses, 95), "ms to first audio"),
    ("First danger alert (avg)", stat(vals("firstDangerS"), statistics.mean), "s after first caller line"),
    ("Firestore write (avg)", stat(vals("writeAvgMs"), statistics.mean), "ms per field update"),
    ("Case summary (avg)", stat(post("consolidateMs"), statistics.mean), "ms after hang-up"),
    ("Tokens per call (avg)", stat(vals("totalTokens"), statistics.mean), "Live API"),
    ("Retries / errors", f"{sum(r['attempts'] - 1 for r in results)} / {sum(1 for r in results if r.get('error'))}", "transient retries / failed calls"),
]
for i, (label, value, unit) in enumerate(kpis):
    col, row = 2 + (i % 4) * 3, 6 + (i // 4) * 4
    for rr in range(row, row + 3):
        for cc in range(col, col + 2):
            ws.cell(row=rr, column=cc).fill = CARD_FILL
    ws.cell(row=row, column=col, value=label).font = Font(color=MUTED, size=9)
    v = ws.cell(row=row + 1, column=col, value=value if value is not None else "—")
    v.font = Font(size=18, bold=True, color=BRAND if i < 4 else INK)
    v.alignment = Alignment(horizontal="left")
    ws.cell(row=row + 2, column=col, value=unit).font = Font(color=MUTED, size=9)
for c in range(2, 15):
    ws.column_dimensions[get_column_letter(c)].width = 13

trow = 19
ws.cell(row=trow, column=2, value="Performance targets").font = Font(bold=True, size=12, color=INK)
for i, h in enumerate(["Metric", "", "", "Target", "Measured", "Stat", "Result"]):
    c = ws.cell(row=trow + 1, column=2 + i, value=h)
    c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
for i, (label, unit, target, measured, st) in enumerate(TARGETS, trow + 2):
    ws.cell(row=i, column=2, value=label)
    ws.cell(row=i, column=5, value=f"≤ {target} {unit}")
    ws.cell(row=i, column=6, value=f"{measured} {unit}" if measured is not None else "—")
    ws.cell(row=i, column=7, value=st)
    ws.cell(row=i, column=8, value="—" if measured is None else "MET" if measured <= target else "MISSED")
color_words(ws, "H", trow + 2, trow + 1 + len(TARGETS))

# Chart data (below the charts)
base = 90
ws.cell(row=base - 1, column=2, value="Chart data").font = Font(bold=True, color=MUTED)
for j, h in enumerate(["Scenario", "Judged pass %", "Strict pass %", "Avg response ms", "First alert s"], 2):
    ws.cell(row=base, column=j, value=h)
for i, sid in enumerate(scen_ids, base + 1):
    rs = by_scen[sid]
    cs = [c for r in rs for c in r["checks"]]
    ws.cell(row=i, column=2, value=sid)
    ws.cell(row=i, column=3, value=round(100 * sum(bool(c["semantic"]) for c in cs) / max(1, len(cs))))
    ws.cell(row=i, column=4, value=round(100 * sum(c["pass"] for c in cs) / max(1, len(cs))))
    ws.cell(row=i, column=5, value=stat(vals("avgResponseMs", rs=rs), statistics.mean))
    ws.cell(row=i, column=6, value=stat([r["firstAlertS"] for r in rs if num(r.get("firstAlertS")) is not None], statistics.mean))
end = base + len(scen_ids)

counts = defaultdict(int)
for r in results:
    counts[status(r, "semantic")] += 1
srow = end + 3
ws.cell(row=srow, column=2, value="Outcome")
ws.cell(row=srow, column=3, value="Calls")
for i, (k, v) in enumerate(sorted(counts.items()), srow + 1):
    ws.cell(row=i, column=2, value=k)
    ws.cell(row=i, column=3, value=v)

groups = defaultdict(lambda: [0, 0])
for r in results:
    groups[r["group"]][0] += sum(bool(c["semantic"]) for c in r["checks"])
    groups[r["group"]][1] += len(r["checks"])
grow = srow + len(counts) + 3
ws.cell(row=grow, column=2, value="Feature group")
ws.cell(row=grow, column=3, value="Pass rate % (judged)")
for i, (k, (ok, n)) in enumerate(groups.items(), grow + 1):
    ws.cell(row=i, column=2, value=k)
    ws.cell(row=i, column=3, value=round(100 * ok / max(1, n)))

buckets = [(0, 1000), (1000, 1500), (1500, 2000), (2000, 2500), (2500, 3500), (3500, 10 ** 9)]
hrow = grow + len(groups) + 3
ws.cell(row=hrow, column=2, value="Response time")
ws.cell(row=hrow, column=3, value="Turns")
for i, (a, b) in enumerate(buckets, hrow + 1):
    ws.cell(row=i, column=2, value=f"{a / 1000:.1f}–{b / 1000:.1f} s" if b < 10 ** 9 else f"> {a / 1000:.1f} s")
    ws.cell(row=i, column=3, value=sum(a <= x < b for x in responses))


def bar(title, cols, anchor, ytitle=None, horizontal=True, first=base, last=end, w=17, legend=True):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title, ch.height, ch.width = title, 9, w
    ch.y_axis.title = ytitle
    for c in cols:
        ch.add_data(Reference(ws, min_col=c, min_row=first, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=first + 1, max_row=last))
    if legend:
        ch.legend.position = "b"
    else:
        ch.legend = None
    ws.add_chart(ch, anchor)


bar("Pass rate per scenario (%)", [3, 4], "B30")
bar("Mia response latency per scenario (avg ms)", [5], "H30", "ms")
bar("Time to first danger alert (s)", [6], "B49", "s")
bar("Response time distribution (turns)", [3], "H49", horizontal=False, first=hrow, last=hrow + len(buckets), w=12, legend=False)
pie = PieChart()
pie.title, pie.height, pie.width = "Call outcomes (judged)", 9, 9
pie.add_data(Reference(ws, min_col=3, min_row=srow, max_row=srow + len(counts)), titles_from_data=True)
pie.set_categories(Reference(ws, min_col=2, min_row=srow + 1, max_row=srow + len(counts)))
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
ws.add_chart(pie, "B68")
bar("Pass rate by feature group (%)", [3], "G68", first=grow, last=grow + len(groups), w=12, legend=False)

# ---------------------------------------------------------------- Test Cases
ws = wb.create_sheet("Test Cases")
first = {sid: by_scen[sid][0] for sid in scen_ids}
table(ws, ["Test ID", "Scenario", "Title", "Group", "Features covered", "Language", "Mocks / injected media", "Expected", "Checks"],
      [(TC[sid], sid, r["title"], r["group"], r["features"], r["language"], r["mocks"], r.get("expected", ""),
        "; ".join(c["name"] for c in r["checks"])) for sid, r in first.items()],
      [8, 24, 34, 10, 44, 22, 28, 50, 50])

# ---------------------------------------------------------------- Results (one row per call)
ws = wb.create_sheet("Results")
rows = []
for r in results:
    d, m = r["dashboard"], r["metrics"]
    pc = m.get("postCall") or {}
    rows.append((TC[r["id"]], r["id"], r["run"], r["incidentId"], status(r, "semantic"), status(r),
                 f"{sum(bool(c['semantic']) for c in r['checks'])}/{len(r['checks'])}", f"{sum(c['pass'] for c in r['checks'])}/{len(r['checks'])}",
                 d.get("callState"), d.get("responseStatus"), d["severity"], d["urgency"], "; ".join(d["dangerIndicators"]),
                 d["peopleCount"], d["voiceStress"], d["trackPoints"], d["route"], d["sceneObservations"],
                 pc.get("consolidated"), pc.get("linkedCases"), r["turns"], r["durationS"], r["firstAlertS"], r["attempts"], r.get("error")))
table(ws, ["Test ID", "Scenario", "Run", "Incident ID", "Judged", "Strict", "Judged checks", "Strict checks", "Call state", "Response status",
           "Severity", "Urgency", "Danger tags (dashboard)", "People", "Voice stress", "GPS points", "Safe route", "Scene obs.",
           "Case summary", "Linked cases", "Turns", "Duration s", "First alert s", "Attempts", "Error"],
      rows, [8, 24, 5, 18, 9, 9, 8, 8, 9, 10, 9, 9, 46, 7, 7, 7, 26, 7, 8, 7, 6, 8, 8, 7, 30])
color_words(ws, "E", 2, len(rows) + 1)
color_words(ws, "F", 2, len(rows) + 1)

# ---------------------------------------------------------------- Consistency (per scenario across runs)
ws = wb.create_sheet("Consistency")
rows = []
for sid in scen_ids:
    rs = by_scen[sid]
    cs = [c for r in rs for c in r["checks"]]
    verdicts = defaultdict(set)
    for c in cs:
        verdicts[c["name"]].add(bool(c["semantic"]))
    flaky = sorted(k for k, v in verdicts.items() if len(v) > 1)
    all_pass = all(status(r, "semantic") == "PASS" for r in rs)
    rows.append((TC[sid], sid, len(rs), sum(status(r, "semantic") == "PASS" for r in rs),
                 round(100 * sum(bool(c["semantic"]) for c in cs) / max(1, len(cs))), round(100 * sum(c["pass"] for c in cs) / max(1, len(cs))),
                 stat(vals("avgResponseMs", rs=rs), statistics.mean), stat([r["firstAlertS"] for r in rs if num(r.get("firstAlertS")) is not None], statistics.mean),
                 stat([r["durationS"] for r in rs], statistics.mean), "; ".join(flaky) or "—",
                 "STABLE" if all_pass else "FLAKY" if any(status(r, "semantic") == "PASS" for r in rs) or flaky else "FAILING"))
table(ws, ["Test ID", "Scenario", "Runs", "Runs fully passing", "Judged pass %", "Strict pass %", "Avg response ms", "Avg first alert s",
           "Avg duration s", "Checks that changed between runs", "Verdict"], rows, [8, 24, 6, 9, 9, 9, 10, 10, 9, 50, 10])
ws.conditional_formatting.add(f"K2:K{len(rows) + 1}", CellIsRule(operator="equal", formula=['"STABLE"'], fill=PASS_FILL))
ws.conditional_formatting.add(f"K2:K{len(rows) + 1}", CellIsRule(operator="equal", formula=['"FLAKY"'], fill=WARN_FILL))
ws.conditional_formatting.add(f"K2:K{len(rows) + 1}", CellIsRule(operator="equal", formula=['"FAILING"'], fill=FAIL_FILL))

# ---------------------------------------------------------------- Detection (per fact)
ws = wb.create_sheet("Detection")
cat = defaultdict(lambda: {"n": 0, "strict": 0, "judged": 0, "at": []})
for r in results:
    for c in r["checks"]:
        if c["name"].startswith(("detects ", "scene: ")) or c["name"] in ("coercion flagged", "plate captured", "address captured", "route guidance given"):
            k = cat[c["name"]]
            k["n"] += 1
            k["strict"] += c["pass"]
            k["judged"] += bool(c["semantic"])
            if num(c.get("atS")) is not None:
                k["at"].append(c["atS"])
rows = [(k, v["n"], v["strict"], v["judged"], round(100 * v["judged"] / max(1, v["n"])), stat(v["at"], statistics.mean), stat(v["at"], statistics.median))
        for k, v in sorted(cat.items(), key=lambda kv: kv[1]["judged"] / max(1, kv[1]["n"]))]
table(ws, ["Fact to detect", "Times expected", "Detected (strict)", "Detected (judged)", "Recall % (judged)", "Avg time to detect s", "Median s"],
      rows, [26, 10, 10, 10, 10, 12, 10])

# ---------------------------------------------------------------- Checks
ws = wb.create_sheet("Checks")
rows = [(TC[r["id"]], r["id"], r["run"], r["incidentId"], c["name"], "PASS" if c["pass"] else "FAIL",
         "PASS" if c["semantic"] else "FAIL", c.get("atS"), c["detail"], c.get("judge", ""))
        for r in results for c in r["checks"]]
table(ws, ["Test ID", "Scenario", "Run", "Incident ID", "Check", "Strict", "Judged", "Detected at s", "Evidence (Mia's tool calls)", "Judge's reason"],
      rows, [8, 24, 5, 18, 26, 7, 7, 8, 70, 50])
color_words(ws, "F", 2, len(rows) + 1)
color_words(ws, "G", 2, len(rows) + 1)

# ---------------------------------------------------------------- Latency (per turn)
ws = wb.create_sheet("Latency")
rows = [(TC[r["id"]], r["id"], r["run"], r["incidentId"], "greeting" if t["turn"] == 0 else t["turn"], round(t["sentAt"] / 1000, 1),
         t["firstAudioMs"], t["completeMs"], t["words"]) for r in results for t in r["turnTimings"]]
table(ws, ["Test ID", "Scenario", "Run", "Incident ID", "Turn", "Sent at s", "First audio ms", "Turn complete ms", "Mia words"],
      rows, [8, 24, 5, 18, 9, 9, 13, 15, 10])
if rows:
    ws.conditional_formatting.add(f"G2:G{len(rows) + 1}", CellIsRule(operator="greaterThan", formula=["2500"], fill=FAIL_FILL))
    ws.conditional_formatting.add(f"G2:G{len(rows) + 1}", CellIsRule(operator="between", formula=["1500", "2500"], fill=WARN_FILL))

# ---------------------------------------------------------------- Performance (per call)
ws = wb.create_sheet("Performance")
M = [("incidentCreateMs", "Incident create ms"), ("connectMs", "Connect ms"), ("greetingFirstAudioMs", "Greeting first audio ms"),
     ("avgResponseMs", "Avg response ms"), ("p50ResponseMs", "p50 response ms"), ("p95ResponseMs", "p95 response ms"),
     ("maxResponseMs", "Max response ms"), ("firstToolS", "First tool call s"), ("firstDangerS", "First danger alert s"),
     ("addressS", "Address confirmed s"), ("routeGuidanceMs", "Route guidance ms"), ("writeAvgMs", "Firestore write avg ms"),
     ("writeP95Ms", "Firestore write p95 ms"), ("writes", "Writes"), ("writeErrors", "Write errors"), ("toolCalls", "Tool calls"),
     ("promptTokens", "Prompt tokens"), ("responseTokens", "Response tokens"), ("totalTokens", "Total tokens"),
     ("framesSent", "Frames sent"), ("audioSentS", "Audio sent s"), ("spokenChunks", "Spoken audio chunks")]
P = [("consolidateMs", "Case summary ms"), ("groundedMs", "Grounded context ms"), ("correlateMs", "Case linking ms")]
rows = [(TC[r["id"]], r["id"], r["run"], r["incidentId"]) + tuple(r["metrics"].get(k) for k, _ in M)
        + tuple((r["metrics"].get("postCall") or {}).get(k) for k, _ in P) for r in results]
summary = [(name, "", "", "") + tuple(stat(vals(k), fn) for k, _ in M) + tuple(stat(post(k), fn) for k, _ in P)
           for name, fn in (("Average", statistics.mean), ("Median", statistics.median), ("Min", min), ("Max", max))]
table(ws, ["Test ID", "Scenario", "Run", "Incident ID"] + [h for _, h in M] + [h for _, h in P],
      rows + [("",) * (4 + len(M) + len(P))] + summary, [8, 24, 5, 18] + [11] * (len(M) + len(P)))
for rr in range(len(rows) + 3, len(rows) + 3 + len(summary)):
    for cc in range(1, 5 + len(M) + len(P)):
        ws.cell(row=rr, column=cc).font = Font(bold=True)
        ws.cell(row=rr, column=cc).fill = CARD_FILL

# ---------------------------------------------------------------- Findings
ws = wb.create_sheet("Findings")
findings = []
manual = HERE / "results" / "findings_manual.json"
if manual.exists():
    for f in json.loads(manual.read_text(encoding="utf-8")):
        findings.append((f.get("test", ""), f.get("scenario", ""), f.get("incidentId", ""), f.get("severity", ""), f.get("area", "Product"),
                         f.get("finding", ""), f.get("evidence", ""), f.get("action", "")))
for r in results:
    tid = TC[r["id"]]
    if r.get("error"):
        findings.append((tid, f"{r['id']}#{r['run']}", r["incidentId"], "High", "Run", "Call failed", r["error"][:300], "Investigate and re-run"))
    for c in r["checks"]:
        if not c["semantic"]:
            sev = "High" if any(k in c["name"] for k in ("detects", "cover kept", "silent", "stay", "incident", "scene")) else "Medium"
            findings.append((tid, f"{r['id']}#{r['run']}", r["incidentId"], sev, "Auto", f"Failed: {c['name']}", c["detail"], c.get("judge", "")))
    slow = [t for t in r["turnTimings"] if t["turn"] > 0 and num(t["firstAudioMs"]) and t["firstAudioMs"] > 2500]
    if slow:
        findings.append((tid, f"{r['id']}#{r['run']}", r["incidentId"], "Low", "Auto", "Slow responses",
                         f"{len(slow)} turn(s) over 2.5 s to first audio (max {max(t['firstAudioMs'] for t in slow)} ms)", ""))
table(ws, ["Test ID", "Scenario#run", "Incident ID", "Severity", "Source", "Finding", "Evidence", "Suggested action / judge"],
      findings or [("", "", "", "", "", "No findings", "", "")], [8, 26, 18, 9, 9, 40, 70, 50])
ws.conditional_formatting.add(f"D2:D{len(findings) + 1}", CellIsRule(operator="equal", formula=['"High"'], fill=FAIL_FILL))
ws.conditional_formatting.add(f"D2:D{len(findings) + 1}", CellIsRule(operator="equal", formula=['"Medium"'], fill=WARN_FILL))

# ---------------------------------------------------------------- Transcripts
ws = wb.create_sheet("Transcripts")
rows = []
for r in results:
    ev = sorted([(l["at"], l["who"].upper(), l["text"]) for l in r["transcript"]]
                + [(t["at"], "TOOL", f"{t['name']} {json.dumps(t['args'], ensure_ascii=False)}") for t in r["tools"]], key=lambda e: e[0])
    rows += [(TC[r["id"]], r["id"], r["run"], r["incidentId"], round(at / 1000, 1), who, text) for at, who, text in ev]
table(ws, ["Test ID", "Scenario", "Run", "Incident ID", "t (s)", "Who", "Text / tool call"], rows, [8, 24, 5, 18, 7, 8, 120])

wb.save(OUT)
print(f"Wrote {OUT}")
